from __future__ import annotations

import os
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Iterator

from langchain_core.document_loaders import BaseBlobParser, Blob
from langchain_core.documents import Document


# python-pptx can only open OOXML (.pptx).
# ODP and legacy .ppt must be converted to PPTX via LibreOffice first.
_NEEDS_CONVERSION: frozenset[str] = frozenset({".odp", ".ppt"})


def _libreoffice_to_pptx(source_path: str) -> tuple[str, str]:
    """Convert *source_path* to PPTX using LibreOffice headless.

    Returns (pptx_path, tmpdir_to_cleanup).  Caller must delete tmpdir.
    """
    binary = shutil.which("libreoffice") or shutil.which("soffice")
    if not binary:
        raise RuntimeError(
            "LibreOffice / soffice not found in PATH. "
            "Install it (e.g. `apt-get install libreoffice`) to ingest ODP / PPT files."
        )
    src = Path(source_path)
    outdir = Path(tempfile.mkdtemp(prefix="ppt-convert-"))
    cmd = [binary, "--headless", "--convert-to", "pptx", "--outdir", str(outdir), str(src)]
    proc = subprocess.run(cmd, capture_output=True, text=True)
    if proc.returncode != 0:
        shutil.rmtree(outdir, ignore_errors=True)
        raise RuntimeError(
            f"LibreOffice conversion failed (exit {proc.returncode}): "
            f"stdout={proc.stdout!r} stderr={proc.stderr!r}"
        )
    converted = outdir / f"{src.stem}.pptx"
    if not converted.exists():
        present = ", ".join(p.name for p in outdir.iterdir()) if outdir.exists() else "<none>"
        shutil.rmtree(outdir, ignore_errors=True)
        raise RuntimeError(
            f"Expected {converted.name} after LibreOffice conversion, found: {present}"
        )
    return str(converted), str(outdir)


class PowerPointParser(BaseBlobParser):
    """Extract text from PPTX / PPT / ODP using python-pptx.

    * PPTX  → opened directly.
    * PPT / ODP → converted to PPTX via LibreOffice, then opened.
    """

    def lazy_parse(self, blob: Blob) -> Iterator[Document]:
        try:
            from pptx import Presentation  # python-pptx
        except ImportError as exc:  # pragma: no cover
            raise ImportError(
                "python-pptx is required. Add `python-pptx>=1.0.0` to requirements.txt."
            ) from exc

        suffix = os.path.splitext(blob.source or "")[1].lower() or ".pptx"
        orig_tmp: str | None = None
        conv_dir: str | None = None

        try:
            # Write the blob to a named temp file with the original extension
            with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
                orig_tmp = tmp.name
                tmp.write(blob.as_bytes())

            # For formats python-pptx cannot read, convert via LibreOffice first
            if suffix in _NEEDS_CONVERSION:
                pptx_path, conv_dir = _libreoffice_to_pptx(orig_tmp)
            else:
                pptx_path = orig_tmp

            prs = Presentation(pptx_path)
            for slide_num, slide in enumerate(prs.slides, start=1):
                texts: list[str] = []
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for para in shape.text_frame.paragraphs:
                            line = " ".join(run.text for run in para.runs).strip()
                            if line:
                                texts.append(line)
                if texts:
                    yield Document(
                        page_content="\n".join(texts),
                        metadata={
                            "source": blob.source or orig_tmp,
                            "slide": slide_num,
                        },
                    )
        finally:
            if orig_tmp and os.path.exists(orig_tmp):
                os.unlink(orig_tmp)
            if conv_dir:
                shutil.rmtree(conv_dir, ignore_errors=True)
